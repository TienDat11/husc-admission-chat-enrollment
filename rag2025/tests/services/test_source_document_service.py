"""Source-document service — ticket 39, criteria 2, 3, 4, 5, 7.

The service is where the registry meets the record store: ``backfill`` derives
entries from what records already cite, ``blind_citations`` names the citations
nothing vouches for, and every write goes down the audit path ticket 12 already
opened (``AuditAction.SOURCE_DOCUMENT_CHANGE`` — no new action value).

Tests run on a synthetic four-version store so the arithmetic is checkable by
hand; the real-store numbers (188 citations / 3 PDFs / 299 corpus pointers) are
a separate test that skips when the gitignored data is absent.
"""
from __future__ import annotations

import json
import sys
from datetime import date, datetime, timezone
from pathlib import Path
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from src.domain.audit import AuditAction  # noqa: E402
from src.domain.canonical import (  # noqa: E402
    AuthorityLevel,
    CanonicalRecord,
    RecordStatus,
    RecordType,
)
from src.domain.source_document import (  # noqa: E402
    SourceDocumentKind,
    classify_document_id,
)
from src.infrastructure.audit_ledger import AuditLedger  # noqa: E402
from src.infrastructure.canonical_record_store import CanonicalRecordStore  # noqa: E402
from src.infrastructure.source_document_registry import (  # noqa: E402
    SourceDocumentRegistry,
)
from src.routers.admin import router as admin_router  # noqa: E402
from src.services.source_document_service import (  # noqa: E402
    DocumentAlreadyRegistered,
    DocumentNotRegisterable,
    DocumentNotRegistered,
    DocumentRetired,
    SourceDocumentService,
)

TTTS_PDF = "new_data/THONG_TIN_TUYEN_SINH_202611.pdf"
HD_PDF = "new_data/2304_BGDDT-GDDH_HD tuyển sinh ĐH, CĐ 2026.pdf"
POINTER = "rag2025/data/chunked/husc_canonical_2026.jsonl#husc_nganh_7220104_2026"

# Ticket 13 identity, same shape as tests/routers/test_admin_lifecycle_api.py:
# one actor per role so a test can name the actor it means.
ACTORS: dict[str, tuple[str, str]] = {
    "admin": ("cb.quan-tri", "admin"),
    "editor": ("cb.bien-tap", "editor"),
    "viewer": ("cb.chi-xem", "viewer"),
}
TOKENS = {role: f"tok-{role}" for role in ACTORS}


def _headers(role: str) -> dict[str, str]:
    return {"X-Admin-Token": TOKENS[role]}


def _record(
    record_id: str, document_id: str | None, page: int | None = 12
) -> CanonicalRecord:
    return CanonicalRecord(
        record_id=record_id,
        record_type=RecordType.DIEM_CHUAN,
        version=1,
        admission_year=2026,
        effective_from=date(2026, 6, 1),
        status=RecordStatus.PUBLISHED,
        authority_level=AuthorityLevel.HUSC_CANONICAL,
        source_document_id=document_id,
        source_url="https://tuyensinh.husc.edu.vn/de-an-2026.pdf",
        source_page=page,
        text="Điểm chuẩn ngành Công nghệ thông tin năm 2026 là 25,5 điểm.",
    )


@pytest.fixture
def repo_root(tmp_path: Path) -> Path:
    """A fake checkout whose cited files exist, so hashing means something."""
    repo = tmp_path / "repo"
    (repo / "new_data").mkdir(parents=True)
    (repo / "new_data" / "THONG_TIN_TUYEN_SINH_202611.pdf").write_bytes(b"%PDF-1.4 ttts")
    (repo / "new_data" / "2304_BGDDT-GDDH_HD tuyển sinh ĐH, CĐ 2026.pdf").write_bytes(
        b"%PDF-1.4 hd"
    )
    return repo


@pytest.fixture
def store(tmp_path: Path) -> CanonicalRecordStore:
    """Four versions: two records cite one PDF, one cites another, one a pointer."""
    store = CanonicalRecordStore(tmp_path / "canonical_store")
    store.write_record(_record("DIEM_CHUAN:NGANH:cntt:2026", TTTS_PDF))
    store.write_record(_record("DIEM_CHUAN:NGANH:luat:2026", TTTS_PDF, page=30))
    store.write_record(_record("CHEINH_SACH:XT:2026", HD_PDF))
    store.write_record(_record("CHI_TIEU:NGANH:7220104:2026", POINTER, page=None))
    return store


@pytest.fixture
def ledger(store: CanonicalRecordStore) -> AuditLedger:
    return AuditLedger(store.root / "audit")


@pytest.fixture
def service(
    store: CanonicalRecordStore, ledger: AuditLedger, repo_root: Path
) -> SourceDocumentService:
    registry = SourceDocumentRegistry(store.root, repo_root)
    return SourceDocumentService(registry, store, ledger)


# -- criterion 2: backfill derives the registry from the record store ----------


def test_backfill_registers_every_cited_document(service: SourceDocumentService) -> None:
    """Nothing is missed: every cited file-backed document ends up registered.

    Four versions, three of which name a file (two cite the same PDF), so the
    report creates two entries — the identity is per *document*, not per
    citation, and the fourth citation is a corpus pointer, skipped by design.
    """
    report = service.backfill(actor_id="system:backfill")
    assert report.scanned_versions == 4
    assert report.citations == 3  # versions naming a file-backed document
    assert sorted(report.registered) == sorted([TTTS_PDF, HD_PDF])
    assert report.already_registered == ()
    assert report.skipped_kinds == {SourceDocumentKind.CHUNK_CORPUS.value: 1}
    assert report.file_absent == ()


def test_backfill_leaves_no_published_record_unregistered(
    service: SourceDocumentService, store: CanonicalRecordStore
) -> None:
    """After backfill, every file-backed citation resolves and is registered.

    ``resolve_provenance`` is the reader's contract (ticket 03); the registry
    landing must not disturb it, and must not leave a published record pointing
    at an id nothing vouches for.
    """
    service.backfill(actor_id="system:backfill")
    for record_id in store.record_ids():
        for version in store.versions(record_id):
            prov = store.resolve_provenance(f"{record_id}@{version}")
            assert prov.source_document_id
            if classify_document_id(prov.source_document_id) is (
                SourceDocumentKind.PDF_DOCUMENT
            ):
                assert service.is_registered(prov.source_document_id)




def test_backfill_is_idempotent_and_writes_no_duplicate_events(
    service: SourceDocumentService, ledger: AuditLedger
) -> None:
    """A second backfill creates nothing and appends nothing to the trail."""
    before = len(ledger.entries())
    first = service.backfill(actor_id="system:backfill")
    second = service.backfill(actor_id="system:backfill")
    assert len(ledger.entries()) == before + len(first.registered)
    assert second.registered == ()
    assert sorted(second.already_registered) == sorted(first.registered)
    assert len(ledger.entries()) == before + len(first.registered)


def test_backfill_records_an_absent_file_as_file_absent(
    tmp_path: Path, store: CanonicalRecordStore, ledger: AuditLedger
) -> None:
    """A citation to a file this machine lacks is still registered.

    Ghi chú 3: the PDFs are gitignored, so a deploy box may not have them. The
    entry records ``file_absent`` as an observation instead of failing the
    backfill and leaving the records unresolvable.
    """
    empty_repo = tmp_path / "no-files"
    empty_repo.mkdir()
    registry = SourceDocumentRegistry(store.root, empty_repo)
    service = SourceDocumentService(registry, store, ledger)
    report = service.backfill(actor_id="system:backfill")
    assert sorted(report.registered) == sorted([TTTS_PDF, HD_PDF])
    assert sorted(report.file_absent) == sorted([TTTS_PDF, HD_PDF])
    entry = registry.get(TTTS_PDF)
    assert entry is not None
    assert entry.file_present is False
    assert entry.sha256 is None
    assert service.is_registered(TTTS_PDF) is True


def test_hash_is_of_the_file_the_registry_points_at(
    service: SourceDocumentService, repo_root: Path
) -> None:
    """The recorded sha256 is the bytes at ``repo_root / document_id``.

    Recomputed here from the resolved path rather than compared to a literal,
    so the assertion catches a hash taken from the wrong file.
    """
    import hashlib

    service.register(TTTS_PDF, actor_id="cb.quan-tri", role="admin")
    entry = service.lookup(TTTS_PDF)
    assert entry is not None
    path = service.resolve_path(TTTS_PDF)
    assert path == (repo_root / TTTS_PDF).resolve()
    assert entry.sha256 == hashlib.sha256(path.read_bytes()).hexdigest()


# -- criterion 3: unresolvable provenance is a named state ---------------------


def test_blind_citations_names_unregistered_published_records(
    service: SourceDocumentService,
) -> None:
    report = service.blind_citations()
    assert report.statuses == ("published",)
    assert report.scanned_versions == 4
    assert report.cited_versions == 4
    assert len(report.items) == 4  # an empty registry blinds every citation
    assert report.counts_by_kind == {
        SourceDocumentKind.PDF_DOCUMENT.value: 3,
        SourceDocumentKind.CHUNK_CORPUS.value: 1,
    }
    assert report.counts_by_document_id[TTTS_PDF] == 2
    keys = {item.record_version_key for item in report.items}
    assert keys == {
        "DIEM_CHUAN:NGANH:cntt:2026@1",
        "DIEM_CHUAN:NGANH:luat:2026@1",
        "CHEINH_SACH:XT:2026@1",
        "CHI_TIEU:NGANH:7220104:2026@1",
    }

    service.backfill(actor_id="system:backfill")
    after = service.blind_citations()
    # Only the corpus pointer remains, by design (Ghi chú 4).
    assert [item.source_document_id for item in after.items] == [POINTER]
    assert after.counts_by_kind == {SourceDocumentKind.CHUNK_CORPUS.value: 1}


def test_blind_citations_never_raises_on_an_unholdable_id(
    service: SourceDocumentService,
) -> None:
    """A ``*.jsonl#slug`` citation is reported, not an exception.

    The real store holds 299 such citations among 487 versions; an existence
    check that delegated straight to a validating ``get`` would crash the
    report (and the HTTP route) before naming a single one.
    """
    assert service.is_registered(POINTER) is False
    assert service.is_registered("../escape.pdf") is False
    assert service.is_registered("/etc/passwd") is False
    assert service.is_registered(None) is False
    assert service.is_registered("   ") is False
    assert service.blind_citations(None).scanned_versions == 4


def test_blind_citations_scans_every_status_when_asked(
    service: SourceDocumentService, store: CanonicalRecordStore
) -> None:
    """``None`` statuses measures the whole store, not just published."""
    store.write_record(
        CanonicalRecord(
            **{
                **_record("CHEINH_SACH:XT:DRAFT:2026", TTTS_PDF).model_dump(),
                "status": RecordStatus.DRAFT,
            }
        )
    )
    assert service.blind_citations().scanned_versions == 4
    assert service.blind_citations(None).scanned_versions == 5


# -- criterion 4: audit rides the existing action ------------------------------


def test_registry_writes_emit_the_existing_audit_action(
    service: SourceDocumentService, ledger: AuditLedger
) -> None:
    """Register / update / retire each append one SOURCE_DOCUMENT_CHANGE.

    No new action value exists — the ticket forbids one outside
    ``src/domain/audit.py:101-116`` — and the hash chain still verifies.
    """
    service.register(TTTS_PDF, actor_id="cb.quan-tri", role="admin", notes="first")
    service.update(
        TTTS_PDF, actor_id="cb.bien-tap", role="editor", notes="re-observed"
    )
    service.retire(
        TTTS_PDF, actor_id="cb.quan-tri", role="admin", reason="superseded in 2027"
    )
    events = [entry.event for entry in ledger.entries()]
    assert {e.action for e in events} == {AuditAction.SOURCE_DOCUMENT_CHANGE}
    # Revisions ride the version keys (AuditEvent has no revision field), in the
    # document id space: document_id@revision, not a record key.
    assert [e.after_version for e in events] == [
        f"{TTTS_PDF}@1",
        f"{TTTS_PDF}@2",
        f"{TTTS_PDF}@3",
    ]
    assert [e.before_version for e in events] == [
        None,
        f"{TTTS_PDF}@1",
        f"{TTTS_PDF}@2",
    ]
    assert {e.entity for e in events} == {"source_document"}
    assert events[2].reason == "superseded in 2027"
    assert ledger.verify().ok


def test_no_registry_action_exists_outside_the_audit_enum(
    service: SourceDocumentService, ledger: AuditLedger
) -> None:
    """Every action the ledger records is a member of the enum."""
    service.register(TTTS_PDF, actor_id="cb.quan-tri", role="admin")
    recorded = {entry.event.action.value for entry in ledger.entries()}
    assert recorded <= {action.value for action in AuditAction}


def test_backfill_appends_one_event_per_created_entry(
    service: SourceDocumentService, ledger: AuditLedger
) -> None:
    """A bulk write is auditable per document, not as one anonymous sweep."""
    report = service.backfill(actor_id="system:backfill")
    events = [entry.event for entry in ledger.entries()]
    assert len(events) == len(report.registered) == 2
    assert {e.target for e in events} == set(report.registered)
    assert {e.actor_id for e in events} == {"system:backfill"}


def test_update_with_nothing_changed_writes_no_event(
    service: SourceDocumentService, ledger: AuditLedger
) -> None:
    """A re-observation that changes nothing is not a decision."""
    service.register(TTTS_PDF, actor_id="cb.quan-tri", role="admin")
    before = len(ledger.entries())
    result = service.update(TTTS_PDF, actor_id="cb.quan-tri", role="admin")
    assert result.event is None
    assert result.entry.revision == 1
    assert len(ledger.entries()) == before


def test_service_without_a_ledger_refuses_nothing_at_this_layer(
    store: CanonicalRecordStore, repo_root: Path
) -> None:
    """The type allows ``ledger=None`` (read wiring); the router refuses writes."""
    registry = SourceDocumentRegistry(store.root, repo_root)
    service = SourceDocumentService(registry, store, None)
    result = service.register(TTTS_PDF, actor_id="cb.quan-tri", role="admin")
    assert result.event is None
    assert service.is_registered(TTTS_PDF) is True


# -- write refusals ------------------------------------------------------------


def test_registering_a_corpus_pointer_is_refused_as_unregisterable(
    service: SourceDocumentService,
) -> None:
    """A real citation of another kind: reportable, never registerable."""
    with pytest.raises(DocumentNotRegisterable):
        service.register(POINTER, actor_id="cb.quan-tri", role="admin")
    with pytest.raises(DocumentNotRegisterable):
        service.update(POINTER, actor_id="cb.quan-tri", role="admin")


def test_registering_an_escaping_id_is_refused(service: SourceDocumentService) -> None:
    with pytest.raises(DocumentNotRegisterable):
        service.register("../outside.pdf", actor_id="cb.quan-tri", role="admin")


def test_register_twice_is_already_registered(service: SourceDocumentService) -> None:
    service.register(TTTS_PDF, actor_id="cb.quan-tri", role="admin")
    with pytest.raises(DocumentAlreadyRegistered):
        service.register(TTTS_PDF, actor_id="cb.quan-tri", role="admin")


def test_update_and_retire_of_an_unknown_id_are_not_registered(
    service: SourceDocumentService,
) -> None:
    with pytest.raises(DocumentNotRegistered):
        service.update("new_data/never.pdf", actor_id="cb.quan-tri", role="admin")
    with pytest.raises(DocumentNotRegistered):
        service.retire(
            "new_data/never.pdf",
            actor_id="cb.quan-tri",
            role="admin",
            reason="gone",
        )


def test_retiring_twice_is_refused(
    service: SourceDocumentService, ledger: AuditLedger
) -> None:
    """Symmetric with update: a second retire would fake a second decision."""
    service.register(TTTS_PDF, actor_id="cb.quan-tri", role="admin")
    service.retire(
        TTTS_PDF, actor_id="cb.quan-tri", role="admin", reason="first retirement"
    )
    before = len(ledger.entries())
    with pytest.raises(DocumentRetired):
        service.retire(
            TTTS_PDF, actor_id="cb.quan-tri", role="admin", reason="again"
        )
    assert len(ledger.entries()) == before


def test_revival_records_the_tombstone_as_before_version(
    service: SourceDocumentService,
) -> None:
    """A re-register reads as "retired at N, registered again at N+1"."""
    service.register(TTTS_PDF, actor_id="cb.quan-tri", role="admin")
    service.retire(
        TTTS_PDF, actor_id="cb.quan-tri", role="admin", reason="wrong file"
    )
    result = service.register(
        TTTS_PDF, actor_id="cb.quan-tri", role="admin", notes="corrected file"
    )
    assert result.entry.revision == 3
    assert result.created is True
    assert result.event is not None
    assert result.event.before_version == f"{TTTS_PDF}@2"


# -- criterion 7: source_page survives ----------------------------------------


def test_source_page_survives_a_registry_landing(
    service: SourceDocumentService, store: CanonicalRecordStore
) -> None:
    """Backfilling the registry does not disturb the page a record cites."""
    service.backfill(actor_id="system:backfill")
    prov = store.resolve_provenance("DIEM_CHUAN:NGANH:luat:2026@1")
    assert prov.source_page == 30
    assert prov.source_document_id == TTTS_PDF


# -- HTTP seam: the endpoints exist and behave --------------------------------


def _client(service: SourceDocumentService) -> TestClient:
    app = FastAPI()
    app.include_router(admin_router)
    app.state.source_document_service = service
    return TestClient(app)


@pytest.fixture
def api(
    service: SourceDocumentService, monkeypatch: pytest.MonkeyPatch
) -> TestClient:
    monkeypatch.setenv(
        "ADMIN_IDENTITIES",
        json.dumps(
            [
                {"token": TOKENS[role], "actor_id": actor, "role": role}
                for role, (actor, _) in ACTORS.items()
            ]
        ),
    )
    return _client(service)


def test_http_register_check_blind_citations_round_trip(
    api: TestClient,
) -> None:
    """Register → check → blind-citations over the mounted admin router."""
    created = api.post(
        "/admin/source-documents/register",
        json={"document_id": TTTS_PDF, "notes": "via http"},
        headers=_headers("admin"),
    )
    assert created.status_code == 201, created.text
    body = created.json()
    assert body["document_id"] == TTTS_PDF
    assert body["sha256"]
    assert body["file_present"] is True

    check = api.get(
        "/admin/source-documents/registered",
        params={"document_id": TTTS_PDF},
        headers=_headers("viewer"),
    )
    assert check.status_code == 200
    assert check.json()["registered"] is True
    assert check.json()["kind"] == "pdf_document"

    blind = api.get(
        "/admin/source-documents/blind-citations",
        params={"scope": "all"},
        headers=_headers("viewer"),
    )
    assert blind.status_code == 200
    report = blind.json()
    assert report["scanned_versions"] == 4
    assert TTTS_PDF not in report["counts_by_document_id"]
    assert report["counts_by_document_id"][POINTER] == 1


def test_http_pointer_is_422_and_unknown_is_404(api: TestClient) -> None:
    refused = api.post(
        "/admin/source-documents/register",
        json={"document_id": POINTER},
        headers=_headers("admin"),
    )
    assert refused.status_code == 422
    missing = api.post(
        "/admin/source-documents/update",
        json={"document_id": "new_data/never.pdf"},
        headers=_headers("admin"),
    )
    assert missing.status_code == 404


def test_http_backfill_and_list(api: TestClient) -> None:
    ran = api.post(
        "/admin/source-documents/backfill", headers=_headers("admin")
    )
    assert ran.status_code == 200, ran.text
    assert sorted(ran.json()["registered"]) == sorted([TTTS_PDF, HD_PDF])

    listed = api.get("/admin/source-documents", headers=_headers("viewer"))
    assert listed.status_code == 200
    view = listed.json()
    assert view["total"] == 2 == view["registered"]
    assert view["counts_by_kind"] == {"pdf_document": 2}


def test_http_retire_requires_a_reason_and_grants_by_role(api: TestClient) -> None:
    api.post(
        "/admin/source-documents/register",
        json={"document_id": TTTS_PDF},
        headers=_headers("admin"),
    )
    no_reason = api.post(
        "/admin/source-documents/retire",
        json={"document_id": TTTS_PDF, "reason": ""},
        headers=_headers("admin"),
    )
    assert no_reason.status_code == 422

    wrong_role = api.post(
        "/admin/source-documents/retire",
        json={"document_id": TTTS_PDF, "reason": "superseded"},
        headers=_headers("viewer"),
    )
    assert wrong_role.status_code == 403

    retired = api.post(
        "/admin/source-documents/retire",
        json={"document_id": TTTS_PDF, "reason": "superseded"},
        headers=_headers("admin"),
    )
    assert retired.status_code == 200
    assert retired.json()["state"] == "retired"


def test_http_write_without_a_ledger_is_503(
    store: CanonicalRecordStore,
    repo_root: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """An unrecorded registry change is the failure mode this ticket removes."""
    monkeypatch.setenv(
        "ADMIN_IDENTITIES",
        json.dumps(
            [{"token": TOKENS["admin"], "actor_id": "cb.quan-tri", "role": "admin"}]
        ),
    )
    registry = SourceDocumentRegistry(store.root, repo_root)
    service = SourceDocumentService(registry, store, None)
    api = _client(service)
    response = api.post(
        "/admin/source-documents/register",
        json={"document_id": TTTS_PDF},
        headers=_headers("admin"),
    )
    assert response.status_code == 503
    assert service.is_registered(TTTS_PDF) is False


def test_lazy_accessor_builds_the_registry_next_to_the_ledger(
    store: CanonicalRecordStore,
    repo_root: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Without an injected service, the router derives ``<store.root>/documents``.

    Same convention as ``<root>/audit/`` and ``<root>/outbox/``, so a store copy
    never arrives without its registry.
    """
    monkeypatch.setenv(
        "ADMIN_IDENTITIES",
        json.dumps(
            [{"token": TOKENS["admin"], "actor_id": "cb.quan-tri", "role": "admin"}]
        ),
    )
    app = FastAPI()
    app.include_router(admin_router)
    app.state.canonical_record_store = store
    app.state.audit_ledger = AuditLedger(store.root / "audit")
    app.state.source_document_repo_root = repo_root
    api = TestClient(app)
    created = api.post(
        "/admin/source-documents/register",
        json={"document_id": TTTS_PDF},
        headers=_headers("admin"),
    )
    assert created.status_code == 201, created.text
    assert (store.root / "documents").is_dir()


def test_lazy_accessor_without_a_store_is_503(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setenv(
        "ADMIN_IDENTITIES",
        json.dumps(
            [{"token": TOKENS["viewer"], "actor_id": "cb.chi-xem", "role": "viewer"}]
        ),
    )
    app = FastAPI()
    app.include_router(admin_router)
    api = TestClient(app)
    response = api.get(
        "/admin/source-documents", headers=_headers("viewer")
    )
    assert response.status_code == 503


# -- T05: a DOCX upload round-trips its headings onto the draft record --------


def _docx_bytes_with_heading() -> bytes:
    import docx

    doc = docx.Document()
    doc.add_paragraph("Phuong thuc xet tuyen", style="Heading 1")
    doc.add_paragraph(
        "Thi sinh dang ky xet tuyen theo phuong thuc hoc ba THPT "
        "voi diem trung binh cac mon dat tu 6.5 tro len trong to hop."
    )
    doc.add_paragraph("Dieu kien xet hoc ba", style="Heading 1")
    doc.add_paragraph(
        "Diem trung binh mon Toan va Ngu van trong ba nam hoc THPT "
        "khong duoi 6.0 diem doi voi moi mon hoc xet tuyen."
    )
    import io

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_upload_docx_round_trips_headings_onto_draft(
    service: SourceDocumentService, store: CanonicalRecordStore
) -> None:
    """The draft an upload creates carries the DOCX headings it was read from.

    The heading text + offsets are what the chunker later turns into the
    ``section`` citation the FE renders, so the round-trip is asserted here
    at the store seam: read the draft back from disk and compare.
    """
    result = service.upload_document(
        _docx_bytes_with_heading(),
        "de-an-tuyen-sinh.docx",
        actor_id="cb.quan-tri",
    )
    assert result.draft_record_id is not None, result.warnings
    draft = store.get_record(result.draft_record_id).latest
    headings = getattr(draft, "docx_headings", None)
    assert headings, "draft must carry the DOCX headings"
    texts = [h["text"] for h in headings]
    assert texts == ["Phuong thuc xet tuyen", "Dieu kien xet hoc ba"]
    assert all(h["offset"] >= 0 for h in headings)
    # Offsets order with the headings they name.
    assert headings[0]["offset"] < headings[1]["offset"]


# -- T17: an XLSX upload drafts one record per data row -----------------------


def _xlsx_bytes(header: list[str], rows: list[list[object]]) -> bytes:
    import io

    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(header)
    for row in rows:
        sheet.append(row)
    buf = io.BytesIO()
    workbook.save(buf)
    return buf.getvalue()


def test_upload_xlsx_writes_one_draft_per_row_with_values(
    service: SourceDocumentService, store: CanonicalRecordStore
) -> None:
    """Each data row becomes a draft carrying that row as its ``values``.

    Three rows, and every draft's payload is asserted against its own input
    row — a bug that duplicated or misassociated rows (e.g. row 3 carrying
    row 1's values) would survive a first-record-only check.
    """
    from src.domain.canonical.semantic_units import structured_payload

    rows = [
        ["7220201", "Công nghệ thông tin", 25.5],
        ["7340101", "Quản trị kinh doanh", 24.75],
        ["7220202", "Kỹ thuật phần mềm", 26.0],
    ]
    data = _xlsx_bytes(["Mã ngành", "Tên ngành", "Điểm chuẩn"], rows)
    result = service.upload_document(data, "diem-chuan.xlsx", actor_id="cb.quan-tri")
    assert result.kind == "xlsx"
    assert len(result.draft_record_ids) == 3
    records = [store.get_record(rid).latest for rid in result.draft_record_ids]
    assert all(r.record_type is RecordType.DIEM_CHUAN for r in records)
    for record, (major_code, major_name, cutoff) in zip(records, rows, strict=True):
        assert dict(structured_payload(record)) == {
            "Mã ngành": major_code,
            "Tên ngành": major_name,
            "Điểm chuẩn": cutoff,
        }
        assert record.text.startswith(f"Mã ngành: {major_code}")


def test_upload_xlsx_unknown_header_falls_back_to_thong_bao(
    service: SourceDocumentService, store: CanonicalRecordStore
) -> None:
    """A sheet naming neither điểm chuẩn nor chỉ tiêu stays THONG_BAO."""
    data = _xlsx_bytes(["Ghi chú", "Giá trị"], [["lịch", "01/07/2026"]])
    result = service.upload_document(data, "ghi-chu.xlsx", actor_id="cb.quan-tri")
    draft = store.get_record(result.draft_record_ids[0]).latest
    assert draft.record_type is RecordType.THONG_BAO
    assert any("THONG_BAO" in warning for warning in result.warnings)


def test_upload_xlsx_reupload_is_already_registered(
    service: SourceDocumentService,
) -> None:
    data = _xlsx_bytes(["Mã ngành", "Điểm chuẩn"], [["7220201", 25.5]])
    service.upload_document(data, "dc.xlsx", actor_id="cb.quan-tri")
    with pytest.raises(DocumentAlreadyRegistered):
        service.upload_document(data, "dc-again.xlsx", actor_id="cb.quan-tri")


def test_detect_kind_recognises_xlsx_and_keeps_docx() -> None:
    """Magic-byte detection: XLSX by its workbook part, DOCX still DOCX."""
    import io

    import docx

    from src.services.magic_bytes import DOCX_KIND, XLSX_KIND, detect_kind

    xlsx = _xlsx_bytes(["A"], [["b"]])
    assert detect_kind(xlsx) == XLSX_KIND

    doc = docx.Document()
    doc.add_paragraph("noi dung van ban")
    doc_buffer = io.BytesIO()
    doc.save(doc_buffer)
    assert detect_kind(doc_buffer.getvalue()) == DOCX_KIND

    assert detect_kind(b"MZ\x90\x00") is None
    assert detect_kind(b"%PDF-1.4") == "pdf_document"


# -- serving policy: get_document_bytes refuses non-corpus locations ---------


def _legacy_entry(registry: SourceDocumentRegistry, document_id: str,
                  repo_root: Path) -> None:
    """Write a registry entry straight to disk, bypassing the classifier.

    Pre-change code registered any repo file (``rag2025/.env``); the new
    classifier refuses that at write time. To exercise the serving policy's
    second line of defence we seed the entry the way the old code did and
    then prove ``get_document_bytes`` refuses to serve it.
    """
    from urllib.parse import quote
    from src.domain.source_document import SourceDocumentKind, SourceDocumentState
    from src.infrastructure.source_document_registry import SourceDocumentEntry

    entry = SourceDocumentEntry(
        document_id=document_id,
        kind=SourceDocumentKind.PDF_DOCUMENT,
        path=str(repo_root / document_id),
        sha256=None,
        file_present=True,
        state=SourceDocumentState.REGISTERED,
        revision=1,
        registered_at=datetime.now(timezone.utc),
        registered_by="cb.quan-tri",
        updated_at=datetime.now(timezone.utc),
        updated_by="cb.quan-tri",
    )
    encoded = quote(document_id, safe="")
    path = registry.documents_dir / f"doc-{encoded}_d.json"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(entry.to_json())


def test_get_document_bytes_refuses_dotfile(
    service: SourceDocumentService, repo_root: Path
) -> None:
    """``rag2025/.env`` is not servable even when a legacy entry exists.

    Registration of a dotfile is now refused by the classifier (OTHER); the
    seeded legacy entry proves the serving path is the second line of
    defence.
    """
    from src.services.source_document_service import SourceDocumentNotServable

    (repo_root / "rag2025").mkdir(parents=True, exist_ok=True)
    (repo_root / "rag2025" / ".env").write_bytes(b"SECRET=1")
    _legacy_entry(service._registry, "rag2025/.env", repo_root)
    with pytest.raises((SourceDocumentNotServable, DocumentNotRegistered)):
        service.get_document_bytes("rag2025/.env")


def test_get_document_bytes_refuses_gitignore(
    service: SourceDocumentService, repo_root: Path
) -> None:
    """A synthetic ``.gitignore`` under ``new_data/`` is not servable."""
    from src.services.source_document_service import SourceDocumentNotServable

    (repo_root / "new_data" / ".gitignore").write_bytes(b"*")
    _legacy_entry(service._registry, "new_data/.gitignore", repo_root)
    with pytest.raises((SourceDocumentNotServable, DocumentNotRegistered)):
        service.get_document_bytes("new_data/.gitignore")


def test_get_document_bytes_serves_new_data_pdf_byte_for_byte(
    service: SourceDocumentService,
) -> None:
    """A legitimate ``new_data/*.pdf`` still serves its exact bytes."""
    service.register(TTTS_PDF, actor_id="cb.quan-tri", role="admin")
    data = service.get_document_bytes(TTTS_PDF)
    assert data == b"%PDF-1.4 ttts"


def test_get_document_bytes_serves_uploads_blob(
    service: SourceDocumentService,
) -> None:
    """An ``uploads/…`` blob written by the uploader is servable."""
    from src.infrastructure.source_document_registry import SOURCE_DOCUMENTS_DIRNAME

    blob_dir = service.registry.root / SOURCE_DOCUMENTS_DIRNAME
    blob_dir.mkdir(parents=True, exist_ok=True)
    blob_path = blob_dir / "deadbeef.pdf"
    blob_path.write_bytes(b"%PDF-1.4 upload-bytes")
    entry = service._registry.create(
        "uploads/deadbeef.pdf",
        sha256="deadbeef",
        file_present=True,
        actor_id="cb.quan-tri",
        now=datetime.now(timezone.utc),
    )
    data = service.get_document_bytes("uploads/deadbeef.pdf")
    assert data == b"%PDF-1.4 upload-bytes"
    _ = entry


# -- defect 5: concurrent identical uploads serialize through the lock --------


def test_two_concurrent_uploads_identical_bytes_yield_one_entry_one_draft(
    service: SourceDocumentService,
    store: CanonicalRecordStore,
    tmp_path: Path,
) -> None:
    """Two simultaneous uploads of the same DOCX bytes: exactly one wins,
    the other gets :class:`DocumentAlreadyRegistered`, the registry holds
    one entry, the store holds one draft record."""
    import threading

    docx_bytes = _docx_bytes_with_heading()
    results: list = []
    errors: list = []

    def _upload():
        try:
            results.append(
                service.upload_document(docx_bytes, "de-an.docx", actor_id="cb.quan-tri")
            )
        except Exception as exc:
            errors.append(exc)

    # Two service instances on the same store root — the real failure mode.
    from src.infrastructure.source_document_registry import SourceDocumentRegistry
    registry2 = SourceDocumentRegistry(store.root, tmp_path / "repo")
    service2 = SourceDocumentService(registry2, store)

    t1 = threading.Thread(target=_upload)

    def _upload2():
        try:
            results.append(
                service2.upload_document(docx_bytes, "de-an.docx", actor_id="cb.quan-tri")
            )
        except Exception as exc:
            errors.append(exc)

    t2 = threading.Thread(target=_upload2)
    t1.start()
    t2.start()
    t1.join()
    t2.join()

    # Exactly one succeeded, one raised DocumentAlreadyRegistered.
    assert len(results) == 1
    assert len(errors) == 1
    assert isinstance(errors[0], DocumentAlreadyRegistered)

    # Registry holds exactly one live entry.
    entry = service.registry.get(results[0].document_id)
    assert entry is not None and entry.is_registered

    # Store holds exactly one draft.
    draft_id = results[0].draft_record_id
    assert draft_id is not None
    assert store.versions(draft_id) == (1,)


# -- defect 6: interrupted commit leaves no stranded draft -------------------


def test_failure_after_draft_write_retries_successfully(
    service: SourceDocumentService,
    store: CanonicalRecordStore,
) -> None:
    """Inject a failure between draft write and registry commit. The retry
    succeeds and does not raise a version conflict: the draft writer reuses
    the already-written v1 record instead of re-appending it."""
    from src.services import source_document_service as svc_module

    docx_bytes = _docx_bytes_with_heading()

    original_create = service._registry.create
    call_count = [0]

    def flaky_create(document_id, *, sha256, file_present, actor_id,
                     notes=None, now=None):
        call_count[0] += 1
        if call_count[0] == 1:
            # Simulate a crash after the draft was written but before the
            # registry entry landed: the draft record already exists in the
            # store, the registry entry does not.
            raise OSError("simulated crash after draft write")
        return original_create(
            document_id,
            sha256=sha256,
            file_present=file_present,
            actor_id=actor_id,
            notes=notes,
            now=now,
        )

    monkeypatch_create = flaky_create
    service._registry.create = monkeypatch_create  # type: ignore[method-assign]

    try:
        # First attempt: draft is written, then create raises.
        with pytest.raises(OSError, match="simulated crash"):
            service.upload_document(docx_bytes, "de-an.docx", actor_id="cb.quan-tri")

        # Verify the draft survived but the registry entry was not created.
        import hashlib
        digest = hashlib.sha256(docx_bytes).hexdigest()
        document_id = f"uploads/{digest}.docx"
        assert service._registry.get(document_id) is None

        # Retry: the draft writer reuses the existing v1, the registry entry
        # is created, and no version conflict is raised.
        result = service.upload_document(
            docx_bytes, "de-an.docx", actor_id="cb.quan-tri"
        )
        assert result.document_id == document_id
        assert result.draft_record_id is not None
        # The draft is at version 1 (reused, not re-appended).
        assert store.versions(result.draft_record_id) == (1,)
    finally:
        service._registry.create = original_create  # type: ignore[method-assign]
